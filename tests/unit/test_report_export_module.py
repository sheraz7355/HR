"""shared/report_export.py helpers that every module's exports rely on."""
import io
from datetime import date

import openpyxl
from PyPDF2 import PdfReader

from shared import report_export as rx


def test_file_name_is_safe_and_keeps_iso_dates(monkeypatch):
    monkeypatch.setattr(rx, "company_name", lambda: "A & B Traders (Pvt) Ltd.")
    name = rx.export_filename("Profit &amp; Loss / by Label", "xlsx",
                              "2026-07-01_to_2027-06-30")
    assert name == "A_and_B_Traders_Pvt_Ltd_Profit_and_Loss_by_Label_2026-07-01_to_2027-06-30.xlsx"
    assert rx.export_filename("Stock", "pdf", date(2026, 10, 5)).endswith("_2026-10-05.pdf")


def test_quantity_columns_get_a_format_without_a_stray_dot():
    fmts = rx._resolve_qty_formats([[1, 10.0], [2, 2.5]], {0: rx.QTY_FMT, 1: rx.QTY_FMT})
    assert fmts == {0: "#,##0", 1: "#,##0.00"}


def test_excel_report_has_the_standard_furniture():
    buf = rx.build_excel("Low Stock", ["SKU", "Qty", "Value"],
                         [["A", 3, 10.5], ["B", 4, 2.25], ["", "", 12.75]],
                         row_kinds=["plain", "plain", "grand"], period="Today",
                         filters=["Warehouse: Main"], col_formats={1: rx.QTY_FMT})
    ws = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active
    assert ws.cell(row=2, column=1).value == "Low Stock"
    assert ws.cell(row=3, column=1).value == "Today"
    assert "Warehouse: Main" in ws.cell(row=4, column=1).value
    assert ws.freeze_panes == "A6"
    assert ws.print_title_rows == "$5:$5"
    # The filter covers the list, not the grand total under it.
    assert ws.auto_filter.ref == "A5:C7"
    assert ws.cell(row=6, column=2).number_format == "#,##0"
    assert ws.cell(row=8, column=3).fill.start_color.rgb.endswith(rx.SLATE)


def test_pdf_escapes_ampersands_and_numbers_pages(monkeypatch):
    monkeypatch.setattr(rx, "company_name", lambda: "Co")
    buf = rx.build_pdf("Parties", ["Name", "Amount"],
                       [["A & B <Traders>", 5.0]] * 90)
    r = PdfReader(io.BytesIO(buf.getvalue()))
    text = r.pages[0].extract_text()
    assert "A & B <Traders>" in text
    assert f"Page 1 of {len(r.pages)}" in text and len(r.pages) > 1


def test_csv_has_bom_and_iso_dates():
    data = rx.build_csv(["Date", "Amount"], [[date(2026, 1, 2), 1.5]]).getvalue()
    assert data.startswith(b"\xef\xbb\xbf")
    assert b"2026-01-02,1.5" in data
