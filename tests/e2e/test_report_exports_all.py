"""Every report exports, and every export carries the same furniture.

Drives each report's export route through the real app (in-process client)
in Excel, PDF and — where offered — CSV, and checks the file rather than
just the status code:

- right mimetype and a descriptive file name (Company_Report_Period.ext);
- Excel: company / title / period / generated-by heading, column headers on
  row 5, frozen below the header, repeated on every printed page, fitted to
  the page width, "Page X of Y" footer, money as numbers;
- PDF: title present, "Page 1 of N" footer, generated line;
- CSV: UTF-8 with BOM, header row first.

Plus the figures that are new in exports: aged receivables buckets (FIFO),
the party ledger's running balance, the stock valuation's GL reconciliation.
"""
import csv
import io
import os
import re
import tempfile
from datetime import datetime, timedelta
from decimal import Decimal

import openpyxl
import pytest
from PyPDF2 import PdfReader

_TMP_DB = os.path.join(tempfile.gettempdir(), "erp_report_exports_all.db")
if os.path.exists(_TMP_DB):
    os.remove(_TMP_DB)
os.environ["DATABASE_URL"] = (os.environ.get("TEST_DATABASE_URL")
                              or "sqlite:///" + _TMP_DB.replace("\\", "/"))

from app import app as flask_app  # noqa: E402
from shared.extensions import db  # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def w():
    c = flask_app.test_client()
    c.get("/")
    from hr_app.models.user import User
    from shared.models.company import CompanyMembership, Company
    from shared.models.ledger import ChartOfAccount, JournalEntry, JournalLine
    from shared.models.exec_report import ExecAccountSelection
    from shared.tenancy import set_current_company, unscoped
    with flask_app.app_context():
        user = (User.query.filter_by(is_super_admin=True).order_by(User.id).first()
                or User.query.order_by(User.id).first())
        uid = user.id
        with unscoped():
            m = (CompanyMembership.query.filter_by(user_id=uid)
                 .order_by(CompanyMembership.id).first())
            cid = m.company_id
        set_current_company(cid)
        from shared.report_export import company_name as _cn
        company_name = _cn()

        # A watched control account with one customer, aged by FIFO:
        # 1,000 posted 45 days ago, 400 posted 5 days ago, 400 received
        # today → 600 left in 31-60 days, 400 current.
        parent = ChartOfAccount(code="9-91-01-01", name="Export Debtors",
                                type="Asset", level=4)
        db.session.add(parent)
        db.session.flush()
        cust = ChartOfAccount(code="9-91-01-01-0001", name="A & B Traders",
                              type="Asset", parent_id=parent.id, level=5)
        db.session.add(cust)
        db.session.flush()
        now = datetime.utcnow()
        for days, dr, cr, num in ((45, 1000, 0, "EXP-1"), (5, 400, 0, "EXP-2"),
                                  (0, 0, 400, "EXP-3")):
            e = JournalEntry(voucher_type="JV", voucher_id=0, voucher_number=num,
                             description="export test", created_by=uid,
                             entry_date=now - timedelta(days=days), is_posted=True)
            db.session.add(e)
            db.session.flush()
            db.session.add(JournalLine(journal_entry_id=e.id, account_id=cust.id,
                                       debit=Decimal(dr), credit=Decimal(cr)))
        db.session.add(ExecAccountSelection(account_id=parent.id,
                                            include_children=True,
                                            created_by=uid))
        db.session.commit()
        cust_id = cust.id
        from inventory_app.models.product import InvProduct
        pid = InvProduct.query.order_by(InvProduct.id).first().id
    with c.session_transaction() as s:
        s["_user_id"] = str(uid)
        s["_fresh"] = True
        s["company_id"] = cid
    return {"c": c, "cust": cust_id, "product": pid, "company": company_name}


def _get(w, url):
    r = w["c"].get(url)
    assert r.status_code == 200, (url, r.status_code, r.data[-800:])
    return r


def _filename(r):
    cd = r.headers.get("Content-Disposition", "")
    m = re.search(r'filename="?([^";]+)"?', cd)
    assert m, cd
    return m.group(1)


def _xlsx(w, url):
    r = _get(w, url)
    assert r.mimetype == XLSX, (url, r.mimetype)
    assert _filename(r).endswith(".xlsx")
    return openpyxl.load_workbook(io.BytesIO(r.data)), r


def _pdf(w, url):
    r = _get(w, url)
    assert r.mimetype == "application/pdf", (url, r.mimetype)
    assert _filename(r).endswith(".pdf")
    return PdfReader(io.BytesIO(r.data)), r


def _assert_report_sheet(ws, title_part):
    assert title_part.lower() in str(ws.cell(row=2, column=1).value).lower()
    assert "Generated" in str(ws.cell(row=4, column=1).value)
    assert ws.freeze_panes is not None, "header must stay put when scrolling"
    assert ws.print_title_rows, "header must repeat on every printed page"
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.page_setup.fitToWidth == 1
    assert "&P" in (ws.oddFooter.right.text or "")


def _assert_report_pdf(reader, title_part):
    first = reader.pages[0].extract_text()
    assert title_part.lower() in first.lower(), first[:300]
    n = len(reader.pages)
    assert f"Page 1 of {n}" in first
    assert "Generated" in first


# Every report export, as (url, title fragment, formats offered).
REPORTS = [
    ("/finance/trial-balance?from=2026-01-01&to=2026-12-31", "Trial Balance", "excel pdf"),
    ("/finance/profit-loss?from=2026-01-01&to=2026-12-31", "Profit", "excel pdf"),
    ("/finance/balance-sheet?to=2026-12-31", "Balance Sheet", "excel pdf"),
    ("/finance/socie", "Equity", "excel pdf"),
    ("/finance/cash-flow", "Cash Flow", "excel pdf"),
    ("/finance/ledger?mode=all&from=2026-01-01&to=2026-12-31", "Ledger", "excel pdf"),
    ("/finance/label-pl", "Label", "excel pdf csv"),
    ("/inventory/reports/valuation", "Stock Valuation", "excel pdf csv"),
    ("/inventory/reports/low-stock", "Low Stock", "excel pdf csv"),
    ("/inventory/vouchers/product-ledger/list", "Sub-Ledgers", "excel pdf csv"),
    ("/executive/receivables", "Aged Receivables", "excel pdf csv"),
    ("/executive/payables", "Aged Payables", "excel pdf csv"),
    ("/executive/integrity", "Books Integrity", "excel pdf csv"),
    ("/fixed-assets/reports/", "Fixed Asset Register", "excel pdf csv"),
    ("/reports/export-attendance?month=10&year=2026", "Attendance", "excel pdf csv"),
    ("/reports/export-leaves?year=2026", "Leave", "excel pdf csv"),
    ("/settings/audit-log.csv", "Audit Log", "excel pdf"),
]


def _with(url, key, val):
    return url + ("&" if "?" in url else "?") + f"{key}={val}"


@pytest.mark.parametrize("url,title,formats", REPORTS)
def test_report_exports_excel(w, url, title, formats):
    if "excel" not in formats:
        pytest.skip("no Excel export")
    wb, r = _xlsx(w, _with(url, "format", "excel"))
    # The company is in the file name, so a folder of exports sorts itself.
    assert _filename(r).startswith(re.sub(r"[^A-Za-z0-9]+", "_", w["company"])[:10])
    ws = wb.worksheets[0]
    if title == "Ledger" and ws.title != "Summary":
        # One account with entries: its own sheet, titled with the account.
        title = ws.title
    _assert_report_sheet(ws, title)


@pytest.mark.parametrize("url,title,formats", REPORTS)
def test_report_exports_pdf(w, url, title, formats):
    reader, _r = _pdf(w, _with(url, "format", "pdf"))
    _assert_report_pdf(reader, title.split()[0] if title != "Ledger" else "Ledger")


@pytest.mark.parametrize("url,title,formats", [r for r in REPORTS if "csv" in r[2]])
def test_report_exports_csv(w, url, title, formats):
    r = _get(w, _with(url, "format", "csv"))
    assert r.mimetype == "text/csv"
    assert r.data.startswith(b"\xef\xbb\xbf"), "BOM so Excel reads UTF-8"
    rows = list(csv.reader(io.StringIO(r.data.decode("utf-8-sig"))))
    assert rows and rows[0], "header row first"


def test_registers_and_voucher_exports(w):
    for url in ("/invoicing/registers/export?fmt=excel&doc=sales",
                "/invoicing/registers/export?fmt=pdf&doc=purchase",
                "/accounting/registers/vouchers/export?fmt=excel"):
        r = _get(w, url)
        assert _filename(r).count("_") >= 2
    r = _get(w, "/invoicing/registers/export?fmt=csv&doc=sales")
    assert r.mimetype == "text/csv"


def test_aged_receivables_split_by_fifo_bucket(w):
    wb, _ = _xlsx(w, "/executive/receivables?format=excel")
    ws = wb.worksheets[0]
    headers = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers[:6] == ["Code", "Party / account", "Group", "Debit",
                           "Credit", "Receivables"]
    assert "31–60 days" in headers and "Avg days" in headers
    row = next(r for r in range(6, ws.max_row + 1)
               if ws.cell(row=r, column=2).value == "A & B Traders")
    vals = dict(zip(headers, [ws.cell(row=row, column=c).value
                              for c in range(1, len(headers) + 1)]))
    assert vals["Receivables"] == 1000
    assert vals["Current"] == 400
    assert vals["31–60 days"] == 600
    assert vals["Oldest days"] == 45
    # Grand total row closes the sheet.
    assert "Total" in str(ws.cell(row=ws.max_row, column=2).value)


def test_party_ledger_export_runs_oldest_first_with_a_balance(w):
    wb, _ = _xlsx(w, f"/executive/party/{w['cust']}?format=excel")
    ws = wb.worksheets[0]
    body = [[ws.cell(row=r, column=c).value for c in range(1, 7)]
            for r in range(6, ws.max_row + 1)]
    balances = [b[5] for b in body[:-1]]
    assert balances == [1000, 1400, 1000]
    assert body[-1][2] == "Balance" and body[-1][5] == 1000


def test_party_names_with_ampersand_render_in_pdf(w):
    reader, _ = _pdf(w, "/executive/receivables?format=pdf")
    assert "A & B Traders" in reader.pages[0].extract_text()


def test_stock_valuation_reconciles_to_the_general_ledger(w):
    wb, _ = _xlsx(w, "/inventory/reports/valuation?format=excel")
    ws = wb.worksheets[0]
    labels = [ws.cell(row=r, column=2).value for r in range(6, ws.max_row + 1)]
    assert "Total" in labels
    assert "Inventory accounts (general ledger)" in labels
    assert any(str(l).startswith("Difference") for l in labels)
    # Quantity columns are whole numbers here, so no "10." stray dot format.
    assert ws.cell(row=6, column=4).number_format == "#,##0"


def test_stock_and_product_ledger_exports(w):
    pid = w["product"]
    wb, _ = _xlsx(w, f"/inventory/reports/stock-ledger?product_id={pid}&format=excel")
    assert wb.worksheets[0].cell(row=5, column=1).value == "Date"
    reader, _ = _pdf(w, f"/inventory/vouchers/product-ledger?product_id={pid}&format=pdf")
    assert "Stock Ledger" in reader.pages[0].extract_text()


def test_general_ledger_workbook_leads_with_a_linked_summary(w):
    wb, _ = _xlsx(w, "/finance/ledger?mode=all&from=2000-01-01&to=2099-12-31&format=excel")
    if len(wb.worksheets) > 2:
        ws = wb.worksheets[0]
        assert ws.title == "Summary"
        assert ws.cell(row=6, column=1).hyperlink is not None


def test_audit_log_csv_default_unchanged(w):
    r = _get(w, "/settings/audit-log.csv")
    assert r.mimetype == "text/csv"
    assert r.data.decode("utf-8", "replace").startswith("When (UTC)")


@pytest.mark.parametrize("page", [
    "/finance/label-pl", "/inventory/reports/valuation",
    "/executive/receivables", "/executive/integrity", "/fixed-assets/reports/",
    "/inventory/vouchers/product-ledger/list"])
def test_report_pages_offer_export_buttons_that_keep_filters(w, page):
    html = _get(w, page + "?as_of=2026-09-30").data.decode()
    assert 'data-export="excel"' in html and 'data-export="pdf"' in html
    # The filter on screen travels with the export link.
    m = re.search(r'href="([^"]*format=excel[^"]*)"', html)
    assert m and "as_of=2026-09-30" in m.group(1).replace("&amp;", "&")


def test_hr_report_builder_export_uses_the_report_workbook(w):
    r = w["c"].post("/reports/export-excel", json={
        "rows": [{"name": "Ali", "present_days": "20", "loan_balance": 1500.5}],
        "columns": ["name", "present_days", "loan_balance"],
        "date_from": "2026-09-01", "date_to": "2026-09-30"})
    assert r.status_code == 200 and r.mimetype == XLSX
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[0]
    assert [ws.cell(row=5, column=c).value for c in (1, 2, 3)] == \
        ["Name", "Present Days", "Loan Balance"]
    assert ws.cell(row=6, column=2).value == 20, "numbers stay numbers"
    assert "2026-09-01" in str(ws.cell(row=3, column=1).value)
